import io
import pytest
import uuid
from uuid import UUID
from reportlab.pdfgen import canvas
from docx import Document as DocxDocument

from app.rag.extractors import DocumentExtractor
from app.rag.chunker import RecursiveTextChunker
from app.rag.embeddings import MockEmbeddingProvider, OpenAIEmbeddingProvider
from app.rag.llm import MockLLMProvider
from app.rag.vector_store import QdrantVectorStore, VectorPoint
from app.rag.rag_service import RAGService
from app.models.project import Project
from app.models.document import Document, DocumentStatus
from app.core.security import create_access_token
from app.services.document_service import DocumentService

def create_sample_pdf_bytes() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Section 1: Enterprise Remote Work Policy. Employees may work remotely up to 3 days weekly.")
    c.showPage()
    c.drawString(100, 750, "Section 2: Home Office Equipment. A one-time stipend of $750 is provided for ergonomic setup.")
    c.showPage()
    c.save()
    return buf.getvalue()

def create_sample_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Employee Benefits Guide", 0)
    doc.add_paragraph("Full-time employees receive 20 days of paid vacation per calendar year.")
    doc.add_paragraph("Comprehensive healthcare coverage begins on day one.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_extract_txt_and_markdown():
    txt_content = b"Simple text file.\nWith multiple lines of policy details."
    extracted = DocumentExtractor.extract(txt_content, "txt")
    assert len(extracted) == 1
    assert "Simple text file." in extracted[0][0]

    md_content = b"# Policy Header\n\n- Point 1: Flexible work\n- Point 2: Security compliance"
    extracted_md = DocumentExtractor.extract(md_content, "md")
    assert len(extracted_md) == 1
    assert "Policy Header" in extracted_md[0][0]

def test_extract_docx():
    docx_bytes = create_sample_docx_bytes()
    extracted = DocumentExtractor.extract(docx_bytes, "docx")
    assert len(extracted) == 1
    assert "Employee Benefits Guide" in extracted[0][0]
    assert "20 days of paid vacation" in extracted[0][0]

def test_extract_pdf():
    pdf_bytes = create_sample_pdf_bytes()
    extracted = DocumentExtractor.extract(pdf_bytes, "pdf")
    assert len(extracted) == 2
    assert "Section 1" in extracted[0][0]
    assert extracted[0][1] == 1  # Page 1
    assert "Section 2" in extracted[1][0]
    assert extracted[1][1] == 2  # Page 2

def test_recursive_text_chunking():
    chunker = RecursiveTextChunker(chunk_size=120, chunk_overlap=20)
    text = (
        "Artificial intelligence and guardrail evaluation are essential for modern LLM applications. "
        "RagGuard continuously inspects retrieved sources to prevent hallucinations and security leaks. "
        "Every retrieved chunk is cross-referenced with enterprise policy baselines before serving."
    )
    pages = [(text, 1)]
    chunks = chunker.chunk_extracted_pages(pages)
    
    assert len(chunks) >= 2
    for i, c in enumerate(chunks):
        assert c.chunk_index == i
        assert c.page_number == 1
        assert "char_count" in c.metadata
        assert c.metadata["chunk_size"] == 120
        assert c.metadata["chunk_overlap"] == 20

@pytest.mark.asyncio
async def test_mock_embedding_provider():
    provider = MockEmbeddingProvider(dimension=1536)
    assert provider.dimension == 1536
    
    query_vec = await provider.embed_query("remote work policy")
    assert len(query_vec) == 1536
    assert isinstance(query_vec[0], float)

    docs_vecs = await provider.embed_documents(["remote work policy", "travel policy"])
    assert len(docs_vecs) == 2
    assert len(docs_vecs[0]) == 1536

@pytest.mark.asyncio
async def test_mock_llm_provider():
    llm = MockLLMProvider()
    prompt = (
        "Context:\n[Source: remote_work_policy.md, Page: 1]\nEmployees are eligible for a $750 home office stipend.\n\n"
        "Question: What is the home office stipend amount?\n\nAnswer:"
    )
    answer = await llm.generate_answer(prompt)
    assert "$750" in answer or "home office stipend" in answer.lower()

@pytest.mark.asyncio
async def test_vector_store_project_isolation():
    store = QdrantVectorStore()
    project_a = uuid.uuid4()
    project_b = uuid.uuid4()

    provider = MockEmbeddingProvider(dimension=1536)
    vec_a = await provider.embed_query("Project A Confidential Info")
    vec_b = await provider.embed_query("Project B Financial Data")

    point_a = VectorPoint(
        id=str(uuid.uuid4()),
        vector=vec_a,
        payload={"project_id": str(project_a), "document_id": str(uuid.uuid4()), "filename": "a.txt", "text": "Secret A"}
    )
    point_b = VectorPoint(
        id=str(uuid.uuid4()),
        vector=vec_b,
        payload={"project_id": str(project_b), "document_id": str(uuid.uuid4()), "filename": "b.txt", "text": "Secret B"}
    )

    await store.upsert_points([point_a, point_b])

    # Search in Project A scope
    results_a = await store.search(query_vector=vec_a, project_id=project_a, top_k=5)
    assert len(results_a) == 1
    assert results_a[0].payload["text"] == "Secret A"
    assert results_a[0].payload["project_id"] == str(project_a)

    # Search in Project B scope
    results_b = await store.search(query_vector=vec_a, project_id=project_b, top_k=5)
    assert len(results_b) == 0 or all(r.payload["project_id"] == str(project_b) for r in results_b)

@pytest.mark.asyncio
async def test_document_upload_and_pipeline(client, test_user, db_session):
    # 1. Create project
    project = Project(name="Upload Test Project", description="Desc", owner_id=test_user.id)
    db_session.add(project)
    await db_session.commit()
    await db_session.refresh(project)

    token = create_access_token(test_user.id)
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Upload markdown file
    file_content = b"# Remote Work Policy\n\nFull time staff get $750 reimbursement for desks."
    files = {
        "file": ("remote_policy.md", file_content, "text/markdown")
    }
    data = {
        "chunk_size": "300",
        "chunk_overlap": "30"
    }

    response = await client.post(
        f"/api/v1/projects/{project.id}/documents/upload",
        files=files,
        data=data,
        headers=headers
    )
    assert response.status_code == 201
    doc_data = response.json()
    assert doc_data["filename"] == "remote_policy.md"
    assert doc_data["file_type"] == "md"
    doc_id = doc_data["id"]

    # 3. Process document pipeline directly for synchronous testing
    await DocumentService.process_document_pipeline(
        document_id=UUID(doc_id),
        file_bytes=file_content,
        db=db_session,
        chunk_size=300,
        chunk_overlap=30
    )

    # 4. Fetch document status
    get_resp = await client.get(f"/api/v1/projects/{project.id}/documents/{doc_id}", headers=headers)
    assert get_resp.status_code == 200
    assert get_resp.json()["status"] == "COMPLETED"
    assert get_resp.json()["chunk_count"] >= 1

    # 5. Fetch chunks
    chunks_resp = await client.get(f"/api/v1/projects/{project.id}/documents/{doc_id}/chunks", headers=headers)
    assert chunks_resp.status_code == 200
    assert len(chunks_resp.json()) >= 1

    # 6. Test RAG query against the uploaded document
    query_payload = {
        "question": "What is the reimbursement amount for desks?",
        "top_k": 3,
        "similarity_threshold": 0.0
    }
    query_resp = await client.post(
        f"/api/v1/projects/{project.id}/query",
        json=query_payload,
        headers=headers
    )
    assert query_resp.status_code == 200
    q_data = query_resp.json()
    assert "answer" in q_data
    assert len(q_data["sources"]) >= 1
    assert q_data["sources"][0]["document"] == "remote_policy.md"
    assert "$750" in q_data["answer"] or "reimbursement" in q_data["answer"].lower()

@pytest.mark.asyncio
async def test_document_upload_unsupported_file(client, test_user, db_session):
    project = Project(name="Invalid File Test", owner_id=test_user.id)
    db_session.add(project)
    await db_session.commit()

    token = create_access_token(test_user.id)
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("malicious.exe", b"binary content", "application/octet-stream")}
    response = await client.post(
        f"/api/v1/projects/{project.id}/documents/upload",
        files=files,
        headers=headers
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]

@pytest.mark.asyncio
async def test_document_upload_unauthorized_user(client, test_user, test_user_2, db_session):
    project = Project(name="Private User 1 Project", owner_id=test_user.id)
    db_session.add(project)
    await db_session.commit()

    token2 = create_access_token(test_user_2.id)
    headers2 = {"Authorization": f"Bearer {token2}"}

    files = {"file": ("notes.txt", b"secret notes", "text/plain")}
    response = await client.post(
        f"/api/v1/projects/{project.id}/documents/upload",
        files=files,
        headers=headers2
    )
    assert response.status_code == 403
